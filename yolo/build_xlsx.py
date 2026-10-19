import json,re,sys
sys.path.insert(0,'.')
from classes import C
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os; os.chdir(os.path.dirname(os.path.abspath(__file__))); B='../build/'
b22=json.load(open(B+'base_el.json')); bnum={e['num']:e for e in b22}
std=[e for e in json.load(open(B+'data.json'))['std'] if e.get('title')!='목차']
fail=[e for e in json.load(open(B+'fail.json')) if e['path'].startswith('별표22') and e['title']!='비고']
# 불합격 행 → 클래스 (원문 대조로 직접 배정). 빈 목록 = 사진 탐지 대상 아님
FM={
'6.1.8.3|승강로 유리':['glass_panel'],
'6.1.9.3|피트 구획/누수':[],'6.3.3|출입문 및 비상문/점검문':['access_door'],
'6.5.1.3|승강로내 돌출물':[],'6.5.1.5|승강로 누수':[],'6.5.2.1|승강로 구획':[],'6.5.2.3|반밀폐식 승강로':[],
'6.5.3.2|승강장문의 문턱 아랫부분':['toe_guard'],
'6.6.4.3.1|카내부 또는 카 지붕의 작업구역':[],'6.6.4.4.1|피트 내부의 작업구역':[],
'6.6.4.6|승강로 외부 작업구역':['access_door'],'6.6.6.1|비상운전 및 작동시험을위한 장치':['emergency_test_panel'],
'7.1|승강장문 및 카문':['landing_door','car_door'],'7.1.1|2개 이상의 카 출입구':[],
'7.3|승강장문 가이드 등':['door_shoe','door_hanger','door_link_rope','vision_panel'],
'7.3|카 문 가이드 등':['door_shoe','vision_panel'],
'7.5.3.4|유리로 된 승강장문, 카문 문짝':['glass_panel'],
'7.6.1|승강장문 및 카문 외부 표면':['landing_door','car_door'],
'7.6.2.2.1|문닫힘안전장치 (수평 개폐식 문)':['door_safety_device'],'7.6.2.3|문닫힘안전장치 (수직 개폐식 문)':['door_safety_device'],
'7.6.3|문닫문의 재-개방':['door_open_button','cop'],
'7.7.2|카 있음 유무 확인 수단':['vision_panel'],'7.7.2.2|카 문의 전망창':['vision_panel'],
'7.9.1|승강장문 잠금장치':['landing_door_lock'],'7.9.1.1|승강장문 잠금':['landing_door_lock'],'7.9.1.8|승강장문 잠금 작용':['landing_door_lock'],
'7.9.2.1|카문 잠금장치':['car_door_lock'],'7.9.3.4|승강장문 도어 클로저':['door_closer'],
'7.9.4.1|승강장문 닫힘 입증 전기안전장치':['landing_door_lock'],'7.10.1|잠금/닫힘 정상 시퀀스의 구성':[],
'7.13.1 · 7.13.2|카 문 닫힘 입증 전기안전장치':['car_door_switch'],'7.15.1|카 정지/카 문의 개방':['door_operator'],
'8.3.1|카의 밀폐식 구획':[],'8.3.2|카의 조립체':[],'8.3.2.1|카 바닥의 기울기':[],'8.3.2.3|카 벽 유리':['glass_panel'],
'8.5.1|에이프런':['apron'],'8.8|카 상부의 설비':['inspection_box','stop_switch'],
'9.1.5 · 10.2.2.1.3|매다는 장치(로프) 파손상태':['suspension_rope'],
'9.2.3.1 · 9.2.3.2|매다는 장치(로프) 끝부분의 고정(소켓)':['rope_termination'],
'9.3|구동기 도르래':['traction_sheave'],'9.5.3|매다는 장치(로프/체인) 이완감지 장치 (포지티브 방식, 간접식)':['slack_rope_switch'],
'10.2.1.1.1|카 추락방지안전장치':['safety_gear'],'10.2.1.4.1|추락방지안전장치 복귀':['safety_gear'],'10.2.1.5|추락방지안전장치 전기적 확인':['safety_gear'],
'10.2.2.1|과속조절기에 의한 작동':['governor'],'10.2.2.1.3|시험 후 과속조절기 로프/관련 부품 상태':['governor_rope','governor'],
'10.2.2.1.6|과속조절기 전기적 확인':['governor'],'10.2.2.3|시험 후 안전 안전로프/관련 부품 상태':[],
'10.3.1|럽처밸브':['rupture_valve'],'10.4.1|유량제한기':[],'10.5.1|멈춤 쇠 장치':['pawl_device'],'10.5.9|멈춤 쇠 장치 전기적 확인':['pawl_device'],
'10.6.1|상승과속 방지장치':['rope_brake'],'10.7.1|개문출발 방지장치':[],'10.7.5|개문출발 방지장치 정지거리':[],
'11.2.1.1|주행안내 레일 및 가이드 슈 등':['guide_rail','guide_shoe'],'12.1.1|완충기':['buffer'],'13.1|구동기':['traction_machine'],
'13.2.2.1.1|브레이크 시스템':['brake'],'13.2.2.2.1|브레이크 제동능력':['brake'],'13.2.2.2.2|브레이크 작동 관련 부품':['brake'],
'13.3.1.1|구동기(유압파워유니트)':['hydraulic_power_unit'],'13.3.2.2.1 · 13.3.2.2.5|잭 설치/고정':['hydraulic_jack'],
'13.3.2.3.2|램 정지수단':['hydraulic_jack'],'13.3.2.6.1|다단 잭 램 정지수단':['hydraulic_jack'],'13.3.3.1.1|배관':['hydraulic_pipe'],
'13.3.5.2.2|체크밸브 작동':[],'13.3.9.2.1|수동-펌프':['hand_pump'],'13.3.10.2|전동기 공회전 방지장치':[],
'14.1.3.1|전기설비의 절연저항':[],'15.2.1.1 · 15.2.4|전기안전장치':[],'16.1.1.4|카의 정상 착상/ 재-착상 정확도':[],
'16.1.2.1|과부하방지장치':['load_weighing_device'],'16.1.2.3|과부하 알림/문 개방':['load_weighing_device'],
'16.1.5.1.1|점검운전 제어':['inspection_box','stop_switch'],'16.1.10|전기적 크리핑방지 시스템':[],
'16.2.1|파이널 리미트 스위치':['limit_switch'],'16.3.3|비상통화장치':['emergency_call'],
}
WHY={'6.1.9.3':'장소 상태(구획·침수) — 물건 아님','6.5.1.3':'장소 상태','6.5.1.5':'장소 상태(누수)','6.5.2.1':'장소 상태(구획)','6.5.2.3':'벽 유무 — 장소 상태',
'6.6.4.3.1':'기계적 장치 외형이 제조사마다 다름 — 사진 모은 뒤 결정','6.6.4.4.1':'위와 같음','7.1.1':'동작 판정','7.10.1':'제어 기능(도어단락 감지)',
'8.3.1':'카 구획 상태','8.3.2':'붕괴 상태','8.3.2.1':'측정 판정','10.2.2.3':'안전로프(유압 간접식) — 사례 드묾, 필요시 추가','10.4.1':'밸브 블록 내부 — 겉모습은 파워유니트',
'10.7.1':'전용 외형 없음(브레이크·제어 기능)','10.7.5':'측정 판정','13.3.5.2.2':'밸브 블록 내부','13.3.10.2':'제어 기능','14.1.3.1':'측정 판정(계측기)',
'15.2.1.1':'회로·동작 판정','16.1.1.4':'동작·측정 판정','16.1.10':'제어 기능'}
keys={c[0] for c in C}
fk=lambda e: e['num']+'|'+e['title']
assert set(FM)=={fk(e) for e in fail}, (set(FM)^{fk(e) for e in fail})
for v in FM.values():
    for k in v: assert k in keys,k
def rx(a): return re.compile(('(?<![가-힣])' if len(a)<=2 else '')+re.escape(a), re.I)
T_b=lambda e:(e.get('title') or '')+' '+e['text']
T_s=lambda e:(e.get('title') or '')+' '+(e.get('gist') or '')+' '+' '.join(p for s in e.get('sec',[]) for p in s['p'])
rows=[]
for i,(key,ko,loc,al,refs) in enumerate(C):
    rs=[rx(a) for a in al]; hit=lambda t:any(r.search(t) for r in rs)
    nb=sum(1 for e in b22 if hit(T_b(e))); ns=sum(1 for e in std if hit(T_s(e)))
    fr=[e for e in fail if key in FM[fk(e)]]
    pr='1' if fr else ('2' if ns>=10 or nb>=10 else '3')
    def rname(r):
        e=bnum[r]; t=(e.get('title') or '').strip()
        if not t or t=='일반사항':
            up=(e.get('path') or '').split(' › ')[-1]; up=up.split(' ',1)[1] if ' ' in up else up
            t=(up+' — '+(t or e['text'][:20].strip()+'…')) if up else e['text'][:24]
        return f'{r} {t}'
    rt=[rname(r) for r in refs]
    rows.append([key,ko,loc,', '.join(al),' / '.join(rt),nb,ns,len(fr),'; '.join(f"{e['num']} {e['title']}" for e in fr),pr])
# 번호 고정: data.yaml 번호 = FIXED 순서(라벨이 이 번호에 묶임). 새 클래스는 FIXED 맨 뒤에 붙일 것.
FIXED='''traction_machine brake traction_sheave suspension_rope rope_termination governor rope_brake access_door
hydraulic_power_unit hydraulic_pipe hand_pump governor_rope guide_rail guide_shoe safety_gear slack_rope_switch
limit_switch hydraulic_jack rupture_valve pawl_device landing_door_lock door_hanger door_link_rope door_closer
toe_guard inspection_box stop_switch door_operator car_door_lock car_door_switch buffer landing_door door_shoe
apron vision_panel load_weighing_device car_door cop door_open_button emergency_call door_safety_device
emergency_test_panel glass_panel pulley sheave_guard suspension_belt controller main_switch manual_handwheel
machine_beam warning_sign shutoff_valve rail_bracket counterweight leveling_device travelling_cable compensation
car_top_balustrade emergency_trap_door counterweight_screen pit_ladder sump unlocking_keyhole door_frame
hall_button fire_switch sill indicator capacity_plate car_handrail light_fixture light_switch socket_outlet
ventilation
tactile_block'''.split()  # 0~73: 1차 확정(2026-10-19), 74~: 추가분
assert len(FIXED)==len(set(FIXED))
# 계산된 우선순위와 달리 정한 것(사용자 지정)
PR={'tactile_block':'2'}  # 점형블록: 불합격 행·언급 수로는 3순위지만 탐지 요청(2026-10-20)
for r in rows:
    if r[0] in PR: r[9]=PR[r[0]]
byk={r[0]:r for r in rows}
gone=[k for k in FIXED if k not in byk]
if gone: print('경고: classes.py 에서 빠진 클래스(번호는 비워 둠):',gone,file=sys.stderr)
# 빠진 클래스는 자리만 남김(번호 당기지 않음)
fixed=[byk.get(k) or [k,'(삭제 — 번호만 유지, 라벨 달지 말 것)','','','',0,0,0,'','-'] for k in FIXED]
# FIXED 에 없는 새 1·2순위는 우선순위·위치 순으로 맨 뒤에(FIXED 에 옮겨 적어 고정할 것)
order=['기계실·구동부','기계실(유압)','승강로','승강로(유압)','승강로(승강장문 뒤)','카 상부·피트','승강장','승강장·카','카','승강장·건물(MRL)','공통']
rest=sorted((r for r in rows if r[0] not in FIXED),key=lambda r:(r[9],order.index(r[2]),r[0]))
new=[r for r in rest if r[9]!='3']
if new: print('경고: FIXED 에 없는 새 클래스(맨 뒤에 붙임, FIXED 에 추가할 것):',[r[0] for r in new],file=sys.stderr)
later=[r for r in rest if r[9]=='3']; rows=fixed+new
wb=Workbook(); ws=wb.active; ws.title='클래스'
H=['id','클래스 키(영문)','한글 이름','위치','별칭·현장용어(라벨러 참고)','근거 조항(별표22, 연혁집 현행)','조문 언급 수','표준화 언급 수','불합격 행 수','관련 불합격 항목(지침 별표2)','우선순위']
ws.append(H)
for i,r in enumerate(rows): ws.append([i]+r)
thin=Side(style='thin',color='DDDDDD')
for c in ws[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='1F4E79'); c.alignment=Alignment(wrap_text=True,vertical='center')
W=[5,22,26,16,34,40,9,9,9,50,8]
for i,w in enumerate(W): ws.column_dimensions[get_column_letter(i+1)].width=w
fills={'1':'FDE8EC','2':'FFF7E0','3':'FFFFFF','-':'F2F2F2'}
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment=Alignment(wrap_text=True,vertical='top'); c.border=Border(bottom=thin)
    row[10].fill=PatternFill('solid',fgColor=fills[str(row[10].value)])
ws.freeze_panes='C2'; ws.auto_filter.ref=ws.dimensions
w4=wb.create_sheet('나중(3순위)')
w4.append(H[1:])
for r in later: w4.append(r)
for i,w in enumerate(W[1:]): w4.column_dimensions[get_column_letter(i+1)].width=w
for c in w4[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='7F7F7F'); c.alignment=Alignment(wrap_text=True,vertical='center')
for row in w4.iter_rows(min_row=2):
    for c in row: c.alignment=Alignment(wrap_text=True,vertical='top'); c.border=Border(bottom=thin)
# 불합격 연결
w2=wb.create_sheet('불합격 행 → 클래스')
w2.append(['조항','검사대상','불합격 내용','클래스','탐지 대상 아님 이유'])
kn={c[0]:c[1] for c in C}
for e in fail:
    ks=FM[fk(e)]
    w2.append([e['num'],e['title'],e['text'],', '.join(f'{kn[k]} ({k})' for k in ks),'' if ks else WHY.get(e['num'],'')])
for i,w in enumerate([14,30,60,40,36]): w2.column_dimensions[get_column_letter(i+1)].width=w
for c in w2[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='1F4E79')
for row in w2.iter_rows(min_row=2):
    for c in row: c.alignment=Alignment(wrap_text=True,vertical='top'); c.border=Border(bottom=thin)
    if not row[3].value: 
        for c in row: c.fill=PatternFill('solid',fgColor='F2F2F2')
w2.freeze_panes='C2'
# 읽는 법
w3=wb.create_sheet('읽는 법')
notes=[
'엘리베이터(별표22) 객체 탐지 클래스 후보 — 검사기준 통합검색 데이터로 뽑음',
'',
f'출처: 승강기 안전기준 연혁집(엘리베이터) 현행 조문 {sum(1 for e in b22 if e["text"].strip())}개, 검사방법 표준화 협의 {len(std)}건(목차 제외), 검사결과 판정 지침 별표2 불합격 내용(별표22) {len(fail)}행, 현장용어 사전.',
'근거 조항: 그 조문(하위 조항 포함) 원문에 별칭 중 하나가 실제로 들어 있는지 스크립트로 확인함.',
f'관련 불합격 항목: 별표2 {len(fail)}행을 하나씩 읽고 직접 배정함(‘불합격 행 → 클래스’ 시트). {len(fail)} 행 중 {sum(1 for e in fail if FM[fk(e)])}행이 클래스와 연결, 사진 탐지 대상 아닌 {sum(1 for e in fail if not FM[fk(e)])}행은 이유를 적고 회색 표시.',
'언급 수: 별칭이 들어 있는 조문·표준화 안건 수(글자 일치, 참고용). 두 글자 이하 별칭은 앞에 한글이 붙지 않을 때만 셈. 일반어(구동기, 잭, 조명 등)는 부풀려질 수 있음.',
'우선순위(제안): 1 = 불합격 항목과 연결된 물건, 2 = 조문 또는 표준화에서 10건 이상 언급, 3 = 나머지. 처음 학습은 1순위부터 권장. 예외(사용자 지정): 점형블록 2순위.',
'',
'라벨링 메모',
'- 상태(파손·녹·누유·미설치)는 클래스에 넣지 않음(요청: 물건 탐지만). 미설치는 ‘그 클래스가 안 잡힘’으로 판단 가능.',
'- 승강장문 잠금장치는 도어스위치와 일체형이 많아 한 클래스로 둠. 따로 볼 필요가 있으면 분리.',
'- 리미트 스위치는 파이널·감속용이 외형이 같아 한 클래스. 구분은 설치 위치로.',
'- 조명기구·콘센트·조명 스위치·환기설비·유리는 위치(기계실·승강로·피트·카) 무관하게 한 클래스. 필요하면 위치 메타데이터로 구분.',
'- 정지장치는 카 상부·피트·기계실 정지스위치 모두. 점검운전 조작반 안의 정지버튼은 조작반 박스와 겹쳐서 따로 박스.',
'- 삼방틀(문틀)과 승강장문(문짝)은 따로 박스. 문턱(씰)은 승강장·카 공통.',
'- 유압 전용 클래스(파워유니트, 잭, 배관, 럽처밸브, 수동펌프, 압력계, 차단밸브, 기름방벽, 멈춤쇠)는 유압식 현장 사진이 있어야 학습 가능.',
'- 1차 학습은 1·2순위 클래스만(‘클래스’ 시트 = data.yaml). 3순위는 ‘나중(3순위)’ 시트에 따로 두었고, 추가할 때는 data.yaml 맨 뒤에 번호를 이어 붙이면 기존 라벨 번호가 안 바뀜.',
'- 클래스 키(영문)는 YOLO data.yaml 의 names 와 같음. id 는 고정(0~73 = 1차 확정 때 우선순위·위치 순서, 74~ = 추가분). 뺀 클래스도 번호는 비워 둠.',
]
for n in notes: w3.append([n])
w3.column_dimensions['A'].width=140
w3['A1'].font=Font(bold=True,size=13)
wb.move_sheet('읽는 법',offset=-wb.sheetnames.index('읽는 법'))
wb.save('엘리베이터_객체인식_클래스_1차.xlsx')
with open('data.yaml','w',encoding='utf-8') as f:
    f.write('# 엘리베이터 객체 탐지 클래스 1차(1·2순위) — 검사기준 통합검색 데이터 기준\n# path/train/val 은 데이터셋에 맞게 채우세요\npath: .\ntrain: images/train\nval: images/val\nnames:\n')
    for i,r in enumerate(rows): f.write(f'  {i}: {r[0]}   # {r[1]}\n')
from collections import Counter
print(len(rows), Counter(r[9] for r in rows), 'fail mapped', sum(1 for e in fail if FM[fk(e)]), 'of', len(fail))

